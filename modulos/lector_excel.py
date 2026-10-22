"""
Paso 1 — Leer la hoja del Excel.

Contrato:
    leer_hoja(ruta, hoja=config.HOJA) -> list[dict]

- Lee desde config.PRIMERA_FILA, columnas config.COL_INI..config.COL_FIN (B..K).
- Cada fila es un dict con las claves de config.COLUMNAS + "fila_excel" (int, número real de fila).
- Devuelve los valores tal cual vienen (sin limpiar).
- Omite filas totalmente vacías.

Estado: LISTO — probado con INVENTARIO_JUNIO_2025 (316 filas).
"""
import openpyxl

import config


def _fila_vacia(valores):
    """True si todos los valores son None o texto en blanco."""
    return all(
        v is None or (isinstance(v, str) and v.strip() == "")
        for v in valores
    )


def leer_hoja(ruta, hoja=config.HOJA):
    """Lee la hoja indicada y devuelve una lista de equipos (dict) sin limpiar."""
    libro = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    try:
        if hoja not in libro.sheetnames:
            raise ValueError(
                f"No existe la hoja '{hoja}' en {ruta}. "
                f"Hojas disponibles: {libro.sheetnames}"
            )
        ws = libro[hoja]

        equipos = []
        filas = ws.iter_rows(
            min_row=config.PRIMERA_FILA,
            min_col=config.COL_INI,
            max_col=config.COL_FIN,
            values_only=True,
        )
        for i, valores in enumerate(filas):
            if _fila_vacia(valores):
                continue
            equipo = dict(zip(config.COLUMNAS, valores))
            equipo["fila_excel"] = config.PRIMERA_FILA + i
            equipos.append(equipo)
        return equipos
    finally:
        libro.close()


if __name__ == "__main__":
    import sys
    equipos = leer_hoja(sys.argv[1])
    print(len(equipos), "filas")
    print(equipos[0])